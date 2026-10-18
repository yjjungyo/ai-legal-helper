# -*- coding: utf-8 -*-
"""
증거 인제스천 및 텍스트 추출 서비스 (Evidence Ingestion & Extraction)

사용자가 업로드한 증거 파일의 무결성을 검증하고(해시, MIME 등), 안전한 로컬 저장소에 보관한 뒤
PDF, DOCX, TXT 파일 등으로부터 텍스트와 위치 정보(EvidenceSpan)를 추출하는 기능을 담당한다.
"""

import os
import hashlib
import uuid
import logging
from datetime import datetime
from typing import List, Tuple, Optional, Any
from pathlib import Path
from ai_legal_helper.utils.datetime_utils import get_kst_now

import filetype

from sqlalchemy.orm import Session
from fastapi import UploadFile, HTTPException

from ai_legal_helper.database.models import EvidenceDocument, EvidenceSpan, EvidenceClaim, Case
from ai_legal_helper.observability.decorators import trace
from ai_legal_helper.observability.context import get_current_span_id
from ai_legal_helper.observability.langfuse_client import update_span

# 보안: 지원하는 MIME 타입 및 확장자 (추가 기능 플래그로 조절 가능)
ALLOWED_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain",
}
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB

logger = logging.getLogger(__name__)

def compute_sha256(file_path: str) -> str:
    """파일의 SHA-256 해시를 계산한다."""
    sha256_hash = hashlib.sha256()
    with open(file_path, "rb") as f:
        for byte_block in iter(lambda: f.read(4096), b""):
            sha256_hash.update(byte_block)
    return sha256_hash.hexdigest()

def get_storage_dir() -> str:
    """안전한 로컬 저장소 경로를 가져오고 존재하지 않으면 생성한다."""
    storage_path = os.environ.get("EVIDENCE_STORAGE_PATH", "./storage/evidence")
    os.makedirs(storage_path, exist_ok=True)
    return storage_path

@trace
async def ingest_evidence(db: Session, case_id: str, file: UploadFile) -> EvidenceDocument:
    """
    업로드된 증거 파일을 검증하고 저장한 뒤 EvidenceDocument 레코드를 생성한다.
    
    1. 파일 크기 및 확장자 검사
    2. UUID 기반 파일명으로 안전하게 저장 (경로 탈출 방지)
    3. 파일 시그니처(MIME) 검사 및 SHA-256 해시 추출
    4. DB에 레코드 생성
    """
    # 1. 대상 Case 존재 여부 확인
    case = db.query(Case).filter_by(id=case_id).first()
    if not case:
        raise HTTPException(status_code=404, detail="해당 사건을 찾을 수 없습니다.")

    # 2. 파일 확장자 검사
    ext = Path(file.filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"허용되지 않는 파일 형식입니다: {ext}")
        
    # 3. 안전한 저장소 위치에 파일 저장
    safe_key = f"{uuid.uuid4()}{ext}"
    storage_dir = get_storage_dir()
    file_path = os.path.join(storage_dir, safe_key)
    
    size_bytes = 0
    try:
        with open(file_path, "wb") as buffer:
            while chunk := await file.read(8192):
                size_bytes += len(chunk)
                if size_bytes > MAX_FILE_SIZE:
                    raise HTTPException(status_code=400, detail="파일 용량이 10MB를 초과했습니다.")
                buffer.write(chunk)
    except Exception as e:
        if os.path.exists(file_path):
            os.remove(file_path)
        raise HTTPException(status_code=500, detail=f"파일 저장 중 오류가 발생했습니다: {str(e)}")
        
    # 4. 파일 무결성 및 형식 검증
    kind = filetype.guess(file_path)
    mime_type = kind.mime if kind else "text/plain" # txt는 filetype에서 안나올 수 있음
    
    if mime_type not in ALLOWED_MIME_TYPES and ext != ".txt":
        os.remove(file_path)
        raise HTTPException(status_code=400, detail=f"유효하지 않은 파일 시그니처입니다: {mime_type}")
        
    file_hash = compute_sha256(file_path)
    
    # 중복 해시 검사 (동일 사건 내에서)
    existing_doc = db.query(EvidenceDocument).filter_by(case_id=case_id, sha256=file_hash).first()
    if existing_doc:
        os.remove(file_path)  # 로컬 저장본 폐기
        raise HTTPException(
            status_code=409, 
            detail=f"이미 동일한 내용의 증거 문서가 등록되어 있습니다. (기존 등록 파일: '{existing_doc.original_filename}')"
        )

    # 5. DB 저장
    doc = EvidenceDocument(
        case_id=case_id,
        original_filename=file.filename or "unknown",
        safe_storage_key=safe_key,
        media_type=mime_type,
        size_bytes=size_bytes,
        sha256=file_hash,
        status="UPLOADED",
        source_kind=ext.replace(".", "").upper()
    )
    
    db.add(doc)
    db.commit()
    db.refresh(doc)
    
    cur_span = get_current_span_id()
    if cur_span:
        try:
            update_span(cur_span, {
                "document_id": doc.id,
                "case_id": case_id,
                "size_bytes": size_bytes,
                "mime_type": mime_type,
            })
        except Exception:
            pass

    return doc

@trace
def extract_text_from_evidence(db: Session, doc_id: str) -> bool:
    """
    EvidenceDocument에 저장된 파일에서 텍스트를 추출하고 EvidenceSpan 레코드들을 생성한다.
    
    - PDF: pdfplumber를 이용하여 페이지별 텍스트 추출
    - DOCX: python-docx를 이용하여 문단별 텍스트 추출
    - TXT: 줄바꿈 단위로 문단 추출
    
    반환값은 추출 성공 여부이다. 성공 시 status가 REVIEW_REQUIRED 등으로 변경된다.
    """
    doc = db.query(EvidenceDocument).filter_by(id=doc_id).first()
    if not doc:
        return False
        
    storage_dir = get_storage_dir()
    file_path = os.path.join(storage_dir, doc.safe_storage_key)
    
    if not os.path.exists(file_path):
        doc.status = "FAILED"
        doc.error_code = "FILE_NOT_FOUND"
        db.commit()
        return False
        
    doc.status = "EXTRACTING"
    db.commit()
    
    try:
        # 재추출 시 이전 spans 및 승인되지 않은(PENDING) claims 중복 방지를 위한 정리
        reviewed_claims_count = (
            db.query(EvidenceClaim)
            .filter_by(evidence_id=doc_id)
            .filter(EvidenceClaim.review_status != "PENDING")
            .count()
        )
        if reviewed_claims_count == 0:
            db.query(EvidenceClaim).filter_by(evidence_id=doc_id).delete(synchronize_session=False)
            db.query(EvidenceSpan).filter_by(evidence_id=doc_id).delete(synchronize_session=False)
            db.flush()

        spans_to_create = []
        ext = Path(doc.original_filename).suffix.lower()
        
        if ext == ".pdf":
            import pdfplumber
            import re
            with pdfplumber.open(file_path) as pdf:
                for i, page in enumerate(pdf.pages):
                    page_text = page.extract_text()
                    if not page_text or not page_text.strip():
                        continue

                    # 1차: 빈 줄(\n\s*\n) 기준 분할
                    raw_paras = [p.strip() for p in re.split(r'\n\s*\n', page_text) if p.strip()]

                    # 빈 줄 구분이 없거나 1개만 나온 경우, 조항(제N조) 또는 번호(1., 2.) 기준으로 세분화
                    if len(raw_paras) <= 1:
                        lines = [line.strip() for line in page_text.splitlines() if line.strip()]
                        chunk = []
                        raw_paras = []
                        for line in lines:
                            if chunk and re.match(r'^(제\s*\d+\s*조|\d+\.|\([0-9가-힣]\)|\[.+\])', line):
                                raw_paras.append("\n".join(chunk))
                                chunk = [line]
                            else:
                                chunk.append(line)
                        if chunk:
                            raw_paras.append("\n".join(chunk))

                    if not raw_paras:
                        raw_paras = [page_text.strip()]

                    for p_idx, para in enumerate(raw_paras):
                        spans_to_create.append(EvidenceSpan(
                            evidence_id=doc.id,
                            page_number=i + 1,
                            paragraph_index=p_idx,
                            text_excerpt=para,
                            extractor_version="pdfplumber-0.11-para",
                            location_quality="EXACT"
                        ))
        elif ext == ".docx":
            import docx
            doc_obj = docx.Document(file_path)
            for i, para in enumerate(doc_obj.paragraphs):
                text = para.text
                if text and text.strip():
                    spans_to_create.append(EvidenceSpan(
                        evidence_id=doc.id,
                        paragraph_index=i,
                        text_excerpt=text.strip(),
                        extractor_version="python-docx-1.2",
                        location_quality="EXACT"
                    ))
        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                lines = f.readlines()
                for i, line in enumerate(lines):
                    if line.strip():
                        spans_to_create.append(EvidenceSpan(
                            evidence_id=doc.id,
                            paragraph_index=i,
                            text_excerpt=line.strip(),
                            extractor_version="txt-basic",
                            location_quality="EXACT"
                        ))
        else:
            doc.status = "EXTRACTION_UNAVAILABLE"
            doc.error_code = "UNSUPPORTED_FORMAT"
            db.commit()
            return False
            
        if not spans_to_create:
            doc.status = "FAILED"
            doc.error_code = "EMPTY_CONTENT"
            db.commit()
            return False
            
        db.add_all(spans_to_create)
        
        doc.status = "REVIEW_REQUIRED" # 향후 AI 추출 시 EXTRACTING 단계로 유지하고 넘길 수 있으나, 일단 텍스트 추출 성공을 의미
        doc.extracted_at = get_kst_now()
        db.commit()
        return True
        
    except Exception as e:
        logger.error(f"Failed to extract text from evidence {doc_id}: {str(e)}")
        doc.status = "FAILED"
        doc.error_code = "EXTRACTION_ERROR"
        db.commit()
        return False

