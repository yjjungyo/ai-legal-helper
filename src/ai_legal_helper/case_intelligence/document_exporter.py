import os
from docx import Document
from datetime import datetime
from ai_legal_helper.database.models import DocumentRevision

def export_docx(revision: DocumentRevision, output_dir: str) -> str:
    """DocumentRevision을 docx 파일로 출력하고 파일 경로를 반환한다."""
    if not os.path.exists(output_dir):
        os.makedirs(output_dir, exist_ok=True)
        
    doc = Document()
    
    # 구조화된 데이터 활용
    content = revision.structured_content or {}
    
    title = content.get("title", "내용증명")
    doc.add_heading(title, 0)
    
    doc.add_paragraph(f"수신인: {content.get('receiver', '미상')}")
    doc.add_paragraph(f"발신인: {content.get('sender', '미상')}")
    
    for section in content.get("body_sections", []):
        for k, v in section.items():
            doc.add_heading(k, level=1)
            doc.add_paragraph(v)
            
    doc.add_heading("결론", level=1)
    doc.add_paragraph(content.get("conclusion", ""))
    
    doc.add_paragraph(f"\n작성일: {datetime.now().strftime('%Y년 %m월 %d일')}")
    
    filename = f"document_{revision.id}_{int(datetime.now().timestamp())}.docx"
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    
    return filepath

def export_pdf(revision: DocumentRevision, output_dir: str) -> str:
    """DocumentRevision을 pdf 파일로 출력하고 파일 경로를 반환한다."""
    if not os.path.exists(output_dir):
        os.makedirs(output_dir, exist_ok=True)
        
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    
    filename = f"document_{revision.id}_{int(datetime.now().timestamp())}.pdf"
    filepath = os.path.join(output_dir, filename)
    
    # 한국어 폰트 설정 (Windows 기준 Malgun Gothic, 없으면 기본 폰트로 시도)
    font_name = "Helvetica"
    malgun_path = "C:\\Windows\\Fonts\\malgun.ttf"
    if os.path.exists(malgun_path):
        pdfmetrics.registerFont(TTFont('Malgun', malgun_path))
        font_name = 'Malgun'
        
    c = canvas.Canvas(filepath, pagesize=A4)
    width, height = A4
    c.setFont(font_name, 12)
    
    content = revision.structured_content or {}
    
    y = height - 50
    title = content.get("title", "내용증명")
    c.setFont(font_name, 16)
    c.drawString(50, y, title)
    y -= 40
    
    c.setFont(font_name, 12)
    c.drawString(50, y, f"수신: {content.get('receiver', '미상')}")
    y -= 20
    c.drawString(50, y, f"발신: {content.get('sender', '미상')}")
    y -= 30
    
    for section in content.get("body_sections", []):
        for k, v in section.items():
            if y < 100:
                c.showPage()
                y = height - 50
                c.setFont(font_name, 12)
            c.setFont(font_name, 14)
            c.drawString(50, y, k)
            y -= 20
            c.setFont(font_name, 12)
            # 간단한 줄바꿈 처리
            lines = v.split('\n')
            for line in lines:
                if y < 50:
                    c.showPage()
                    y = height - 50
                    c.setFont(font_name, 12)
                c.drawString(50, y, line)
                y -= 15
            y -= 15
            
    if y < 100:
        c.showPage()
        y = height - 50
        c.setFont(font_name, 12)
        
    c.setFont(font_name, 14)
    c.drawString(50, y, "결론")
    y -= 20
    c.setFont(font_name, 12)
    c.drawString(50, y, content.get("conclusion", ""))
    y -= 40
    
    c.drawString(50, y, f"작성일: {datetime.now().strftime('%Y년 %m월 %d일')}")
    
    c.save()
    return filepath


